"""
Document processing service for handling various file types.
Extracts text and metadata from PDF, DOCX, CSV, and XLSX files.
"""
import PyPDF2
import pdfplumber
from docx import Document as DocxDocument
import pandas as pd
from typing import Dict, Any, Optional, List
from pathlib import Path
import json
from backend.config import settings
from backend.utils.logger import log
from backend.utils.helpers import sanitize_filename, extract_numbers, extract_dates, detect_language
from backend.models.schemas import DocumentType, DocumentCategory, DocumentMetadata
from backend.services.vector_store import get_vector_store
from anthropic import Anthropic


class DocumentProcessor:
    """Service for processing and analyzing documents."""

    def __init__(self):
        """Initialize the document processor."""
        self.upload_dir = Path(settings.upload_dir)
        self.upload_dir.mkdir(parents=True, exist_ok=True)
        self.vector_store = get_vector_store()
        self.claude_client = Anthropic(api_key=settings.anthropic_api_key)

    def extract_text_from_pdf(self, file_path: Path) -> tuple[str, Dict[str, Any]]:
        """
        Extract text and metadata from a PDF file.

        Args:
            file_path: Path to the PDF file

        Returns:
            Tuple of (extracted text, metadata)
        """
        text = ""
        metadata = {}

        try:
            # Try pdfplumber first (better for tables and complex layouts)
            with pdfplumber.open(file_path) as pdf:
                metadata['page_count'] = len(pdf.pages)

                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n\n"

            # If pdfplumber fails or returns empty, try PyPDF2
            if not text.strip():
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    metadata['page_count'] = len(pdf_reader.pages)

                    # Extract metadata
                    if pdf_reader.metadata:
                        metadata['author'] = pdf_reader.metadata.get('/Author', None)
                        metadata['creation_date'] = pdf_reader.metadata.get('/CreationDate', None)

                    # Extract text
                    for page in pdf_reader.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + "\n\n"

            # Count words
            metadata['word_count'] = len(text.split())

            log.info(f"Extracted {metadata['word_count']} words from PDF with {metadata.get('page_count', 0)} pages")

        except Exception as e:
            log.error(f"Error extracting text from PDF: {e}")
            raise

        return text, metadata

    def extract_text_from_docx(self, file_path: Path) -> tuple[str, Dict[str, Any]]:
        """
        Extract text and metadata from a DOCX file.

        Args:
            file_path: Path to the DOCX file

        Returns:
            Tuple of (extracted text, metadata)
        """
        text = ""
        metadata = {}

        try:
            doc = DocxDocument(file_path)

            # Extract text from paragraphs
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            text = "\n\n".join(paragraphs)

            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells)
                    text += "\n" + row_text

            # Metadata
            core_properties = doc.core_properties
            metadata['author'] = core_properties.author
            metadata['creation_date'] = str(core_properties.created) if core_properties.created else None
            metadata['word_count'] = len(text.split())

            log.info(f"Extracted {metadata['word_count']} words from DOCX")

        except Exception as e:
            log.error(f"Error extracting text from DOCX: {e}")
            raise

        return text, metadata

    def extract_text_from_csv(self, file_path: Path) -> tuple[str, Dict[str, Any]]:
        """
        Extract text and metadata from a CSV file.

        Args:
            file_path: Path to the CSV file

        Returns:
            Tuple of (extracted text, metadata)
        """
        text = ""
        metadata = {}

        try:
            # Read CSV
            df = pd.read_csv(file_path)

            # Convert to text representation
            text = f"CSV File with {len(df)} rows and {len(df.columns)} columns\n\n"
            text += f"Columns: {', '.join(df.columns)}\n\n"

            # Add column statistics
            text += "Data Summary:\n"
            text += df.describe(include='all').to_string()
            text += "\n\n"

            # Add sample rows
            text += "Sample Data (first 10 rows):\n"
            text += df.head(10).to_string()

            # Metadata
            metadata['row_count'] = len(df)
            metadata['column_count'] = len(df.columns)
            metadata['columns'] = df.columns.tolist()

            log.info(f"Extracted data from CSV with {len(df)} rows and {len(df.columns)} columns")

        except Exception as e:
            log.error(f"Error extracting text from CSV: {e}")
            raise

        return text, metadata

    def extract_text_from_xlsx(self, file_path: Path) -> tuple[str, Dict[str, Any]]:
        """
        Extract text and metadata from an XLSX file.

        Args:
            file_path: Path to the XLSX file

        Returns:
            Tuple of (extracted text, metadata)
        """
        text = ""
        metadata = {}

        try:
            # Read all sheets
            excel_file = pd.ExcelFile(file_path)
            sheets = excel_file.sheet_names

            text = f"Excel File with {len(sheets)} sheets\n\n"

            sheet_info = []

            for sheet_name in sheets:
                df = pd.read_excel(file_path, sheet_name=sheet_name)

                text += f"=== Sheet: {sheet_name} ===\n"
                text += f"{len(df)} rows, {len(df.columns)} columns\n"
                text += f"Columns: {', '.join(df.columns)}\n\n"

                # Add statistics
                text += "Data Summary:\n"
                text += df.describe(include='all').to_string()
                text += "\n\n"

                # Add sample rows
                text += "Sample Data (first 5 rows):\n"
                text += df.head(5).to_string()
                text += "\n\n"

                sheet_info.append({
                    'name': sheet_name,
                    'rows': len(df),
                    'columns': len(df.columns)
                })

            # Metadata
            metadata['sheet_count'] = len(sheets)
            metadata['sheets'] = sheet_info

            log.info(f"Extracted data from XLSX with {len(sheets)} sheets")

        except Exception as e:
            log.error(f"Error extracting text from XLSX: {e}")
            raise

        return text, metadata

    def extract_text(self, file_path: Path, file_type: DocumentType) -> tuple[str, Dict[str, Any]]:
        """
        Extract text from a document based on its type.

        Args:
            file_path: Path to the file
            file_type: Type of the document

        Returns:
            Tuple of (extracted text, metadata)
        """
        extractors = {
            DocumentType.PDF: self.extract_text_from_pdf,
            DocumentType.DOCX: self.extract_text_from_docx,
            DocumentType.CSV: self.extract_text_from_csv,
            DocumentType.XLSX: self.extract_text_from_xlsx,
        }

        extractor = extractors.get(file_type)
        if not extractor:
            raise ValueError(f"Unsupported file type: {file_type}")

        return extractor(file_path)

    def analyze_content_with_claude(self, text: str, filename: str) -> Dict[str, Any]:
        """
        Use Claude to analyze document content and extract insights.

        Args:
            text: Document text
            filename: Original filename

        Returns:
            Analysis results
        """
        try:
            # Truncate text if too long (keep first ~50k chars)
            text_sample = text[:50000] if len(text) > 50000 else text

            prompt = f"""Analyse ce document d'entreprise et extrais les informations suivantes:

Document: {filename}

Contenu:
{text_sample}

Fournis une réponse JSON avec:
1. "category": Catégorie principale (finance, strategy, hr, projects, legal, operations, other)
2. "extracted_kpis": Liste des KPIs mentionnés (chiffres clés, métriques)
3. "key_entities": Liste des entités importantes (noms d'entreprises, projets, personnes)
4. "summary": Résumé en 2-3 phrases

Format JSON uniquement."""

            response = self.claude_client.messages.create(
                model=settings.claude_model,
                max_tokens=1024,
                temperature=0.3,
                messages=[{"role": "user", "content": prompt}]
            )

            # Parse response
            content = response.content[0].text

            # Try to extract JSON from response
            import re
            json_match = re.search(r'\{.*\}', content, re.DOTALL)
            if json_match:
                analysis = json.loads(json_match.group())
            else:
                # Fallback
                analysis = {
                    "category": "other",
                    "extracted_kpis": [],
                    "key_entities": [],
                    "summary": "Unable to analyze"
                }

            log.info(f"Claude analysis completed for {filename}")
            return analysis

        except Exception as e:
            log.error(f"Error analyzing content with Claude: {e}")
            return {
                "category": "other",
                "extracted_kpis": [],
                "key_entities": [],
                "summary": ""
            }

    def extract_metadata(self, text: str, base_metadata: Dict[str, Any]) -> DocumentMetadata:
        """
        Extract comprehensive metadata from document text.

        Args:
            text: Document text
            base_metadata: Base metadata from file extraction

        Returns:
            Complete metadata object
        """
        # Extract numbers and dates
        numbers = extract_numbers(text)
        dates = extract_dates(text)

        # Combine with base metadata
        metadata = DocumentMetadata(
            page_count=base_metadata.get('page_count'),
            word_count=base_metadata.get('word_count'),
            creation_date=base_metadata.get('creation_date'),
            author=base_metadata.get('author'),
            extracted_amounts=numbers[:20],  # Limit to first 20
            extracted_dates=dates[:20],
            extracted_kpis=[],
            key_entities=[]
        )

        return metadata

    async def process_document(
        self,
        file_path: Path,
        filename: str,
        file_type: DocumentType,
        document_id: str
    ) -> Dict[str, Any]:
        """
        Process a document: extract text, analyze, and index.

        Args:
            file_path: Path to the file
            filename: Original filename
            file_type: Type of document
            document_id: Unique document ID

        Returns:
            Processing results
        """
        try:
            log.info(f"Processing document: {filename}")

            # Extract text and base metadata
            text, base_metadata = self.extract_text(file_path, file_type)

            if not text.strip():
                raise ValueError("No text could be extracted from document")

            # Analyze with Claude
            claude_analysis = self.analyze_content_with_claude(text, filename)

            # Extract metadata
            metadata = self.extract_metadata(text, base_metadata)
            metadata.extracted_kpis = claude_analysis.get('extracted_kpis', [])
            metadata.key_entities = claude_analysis.get('key_entities', [])

            # Determine category
            category_str = claude_analysis.get('category', 'other')
            try:
                category = DocumentCategory(category_str.lower())
            except ValueError:
                category = DocumentCategory.OTHER

            # Detect language
            language = detect_language(text[:1000])

            # Index in vector store
            vector_metadata = {
                'document_id': document_id,
                'filename': filename,
                'file_type': file_type.value,
                'category': category.value,
                'language': language
            }

            chunks_created = self.vector_store.add_document(
                document_id=document_id,
                text=text,
                metadata=vector_metadata
            )

            log.info(f"Document processed successfully: {filename} ({chunks_created} chunks)")

            return {
                'success': True,
                'category': category,
                'metadata': metadata,
                'chunks_created': chunks_created,
                'summary': claude_analysis.get('summary', '')
            }

        except Exception as e:
            log.error(f"Error processing document {filename}: {e}")
            return {
                'success': False,
                'error': str(e)
            }


# Global instance
_document_processor = None


def get_document_processor() -> DocumentProcessor:
    """Get or create the global document processor instance."""
    global _document_processor
    if _document_processor is None:
        _document_processor = DocumentProcessor()
    return _document_processor
